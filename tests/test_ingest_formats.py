"""PDF / DOCX 解析测试：页偏移、canonical_text、扫描件、加密、页数上限、表格顺序。"""

from __future__ import annotations

import io
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from app.nodes import ingest as ingest_module
from app.nodes.ingest import ingest_document
from app.schemas import ContractReviewState
from tests.fixtures.pdf_builder import make_pdf


def _state(path: Path) -> ContractReviewState:
    return ContractReviewState(review_id="t", source_filename=path.name, source_path=str(path))


def test_pdf_two_pages_offsets(tmp_path: Path) -> None:
    data = make_pdf(
        [
            ["INTERNATIONAL SALES CONTRACT", "ARTICLE 1 GOODS", "Seller: ABC Trading"],
            ["ARTICLE 2 PAYMENT", "Payment by T/T in advance."],
        ]
    )
    path = tmp_path / "contract.pdf"
    path.write_bytes(data)

    update = ingest_document(_state(path))
    assert update.get("status") != "failed"
    assert update["media_type"] == "pdf"
    assert len(update["source_sha256"]) == 64

    pages = update["pages"]
    raw = update["raw_text"]
    assert len(pages) == 2
    assert raw == pages[0].text + "\n\n" + pages[1].text
    for page in pages:
        assert raw[page.start_offset : page.end_offset] == page.text
    assert pages[0].end_offset + 2 == pages[1].start_offset
    assert "ARTICLE 2 PAYMENT" in pages[1].text
    assert any("pypdf" in warning for warning in update["warnings"])


def test_pdf_without_text_layer_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    path.write_bytes(make_pdf([[]]))
    update = ingest_document(_state(path))
    assert update["status"] == "failed"
    assert "扫描件" in update["errors"][0]
    assert update["pages"] == []
    assert update["raw_text"] == ""


def test_pdf_encrypted_is_rejected(tmp_path: Path) -> None:
    reader = PdfReader(io.BytesIO(make_pdf([["ARTICLE 1 GOODS", "Steel coils to be delivered."]])))
    writer = PdfWriter()
    writer.append(reader)
    writer.encrypt(user_password="secret", owner_password="secret")
    buffer = io.BytesIO()
    writer.write(buffer)

    path = tmp_path / "encrypted.pdf"
    path.write_bytes(buffer.getvalue())
    update = ingest_document(_state(path))
    assert update["status"] == "failed"
    assert "加密" in update["errors"][0]


def test_pdf_page_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "two_pages.pdf"
    path.write_bytes(make_pdf([["ARTICLE 1 GOODS"], ["ARTICLE 2 PAYMENT"]]))
    monkeypatch.setattr(
        "app.nodes.ingest.settings",
        replace(ingest_module.settings, max_pages=1),
    )
    update = ingest_document(_state(path))
    assert update["status"] == "failed"
    assert "页数过多" in update["errors"][0]


def test_pdf_invalid_header(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4\nthis is not a real pdf")
    update = ingest_document(_state(path))
    assert update["status"] == "failed"
    assert "PDF" in update["errors"][0]


def test_docx_paragraphs_and_table_in_order(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("INTERNATIONAL SALES CONTRACT")
    document.add_paragraph("ARTICLE 1 GOODS")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Goods"
    table.cell(0, 1).text = "Steel coils"
    table.cell(1, 0).text = "Quantity"
    table.cell(1, 1).text = "5000 tons"
    document.add_paragraph("ARTICLE 2 PAYMENT")
    document.add_paragraph("Payment by L/C at sight.")

    path = tmp_path / "contract.docx"
    document.save(path)
    update = ingest_document(_state(path))

    assert update.get("status") != "failed"
    assert update["media_type"] == "docx"
    assert len(update["pages"]) == 1

    raw = update["raw_text"]
    assert "Goods | Steel coils" in raw
    assert "Quantity | 5000 tons" in raw
    assert raw.index("ARTICLE 1 GOODS") < raw.index("Goods | Steel coils")
    assert raw.index("Goods | Steel coils") < raw.index("ARTICLE 2 PAYMENT")
    page = update["pages"][0]
    assert page.start_offset == 0
    assert page.end_offset == len(raw)


def test_docx_invalid_container(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK\x03\x04not really a docx")
    update = ingest_document(_state(path))
    assert update["status"] == "failed"
    assert "DOCX" in update["errors"][0]


def test_filename_is_sanitized(tmp_path: Path) -> None:
    path = tmp_path / "contract.txt"
    path.write_text("Seller: A\nBuyer: B\n", encoding="utf-8")
    state = _state(path)
    state.source_filename = "../../etc/passwd\n.txt"
    update = ingest_document(state)
    assert update.get("status") != "failed"
    assert "/" not in update["source_filename"]
    assert "\n" not in update["source_filename"]
